"""Word files and scanned pages: what is accepted, how it is read, and what is never marked verified."""

import io
import zipfile

import pymupdf
import pytest

from app.core.config import settings
from app.services import job_runner, ocr_service
from app.services.ocr_service import PageOcr
from sample_data.generate import long_contract_text
from tests.helpers import upload, upload_and_process, upload_version

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# ─── helpers ──────────────────────────────────────────────────────────────────

def make_docx(text: str = None, table=None, extra_zip=None) -> bytes:
    import docx
    d = docx.Document()
    for para in (text or long_contract_text()).split("\n"):
        if para.strip():
            d.add_paragraph(para)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    buf = io.BytesIO()
    d.save(buf)
    data = buf.getvalue()
    if extra_zip:
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
            for item in src.infolist():
                dst.writestr(item, src.read(item.filename))
            for name, content in extra_zip.items():
                dst.writestr(name, content)
        data = out.getvalue()
    return data


def to_scan(pdf_bytes: bytes) -> bytes:
    """An image-only copy of a PDF: no text layer, like a scanner's output."""
    src = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    scan = pymupdf.open()
    for page in src:
        png = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False).tobytes("png")
        new = scan.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, stream=png)
    return scan.tobytes()


def page_texts(pdf_bytes: bytes):
    src = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    return {i + 1: p.get_text().strip() for i, p in enumerate(src)}


async def upload_docx(client, data: bytes, name: str = "Vendor Agreement.docx"):
    return await client.post("/api/contracts/upload", files={"file": (name, data, DOCX_TYPE)})


# ─── Word files ───────────────────────────────────────────────────────────────

async def test_a_word_document_is_converted_analysed_and_verifiable(client, llm):
    data = make_docx()
    r = await upload_docx(client, data)
    assert r.status_code == 202, r.text
    await job_runner.wait_idle()
    d = (await client.get(f"/api/contracts/{r.json()['id']}")).json()
    assert d["status"] == "READY" and d["extraction_status"] == "complete" and d["page_count"] >= 1
    assert d["title"] == "Vendor Agreement" and d["extraction_meta"]["source"]["kind"] == "docx"
    exp = next(f for f in d["fields"] if f["field_key"] == "expiration_date")
    assert exp["display_value"] == "2028-12-31" and exp["status"] == "verified" and exp["page"]
    hl = await client.get(f"/api/contracts/{d['id']}/pages/{exp['page']}/highlights", params={"q": exp["source_quote"]})
    assert hl.json()["matched"] is True                        # the converted page can be highlighted
    flags = (await client.get(f"/api/contracts/{d['id']}/flags")).json()
    conv = [f for f in flags if f["code"] == "CONVERTED_DOCUMENT"]
    assert conv and conv[0]["severity"] == "LOW" and "converted" in conv[0]["reason"].lower()


async def test_the_converted_pdf_and_the_original_word_file_can_both_be_downloaded(client, llm):
    data = make_docx()
    cid = (await upload_docx(client, data, "My Contract.docx")).json()["id"]
    await job_runner.wait_idle()
    pdf = await client.get(f"/api/contracts/{cid}/file")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf" and pdf.content.startswith(b"%PDF-")
    assert "My%20Contract.pdf" in pdf.headers["content-disposition"]
    orig = await client.get(f"/api/contracts/{cid}/file?original=true")
    assert orig.status_code == 200 and orig.headers["content-type"] == DOCX_TYPE and orig.content == data
    assert "My%20Contract.docx" in orig.headers["content-disposition"]


async def test_table_text_is_kept_and_markup_in_the_text_is_not_interpreted(client, llm):
    data = make_docx("1. Fees\nThe fee is <b>bold</b> & more.", table=[["Item", "Amount"], ["Monthly fee", "$12,500"]])
    cid = (await upload_docx(client, data)).json()["id"]
    await job_runner.wait_idle()
    text = " ".join(p["text"] for p in (await client.get(f"/api/contracts/{cid}/text")).json()["pages"])
    assert "$12,500" in text and "Monthly fee" in text and "<b>bold</b>" in text and "& more" in text


async def test_a_word_version_can_be_added_to_a_contract(client, llm, long_pdf):
    d = await upload_and_process(client, long_pdf)
    v = await upload_version(client, d["id"], make_docx(long_contract_text().replace("ninety (90)", "sixty (60)")), "v2.docx", label="v2 (Word)")
    assert v["filename"] == "v2.docx" if "filename" in v else True
    detail = (await client.get(f"/api/contracts/{d['id']}")).json()
    assert [x["label"] for x in detail["versions"]] == ["v1", "v2 (Word)"] and detail["status"] == "READY"


async def test_permanent_delete_removes_the_kept_original_too(client, llm):
    cid = (await upload_docx(client, make_docx())).json()["id"]
    await job_runner.wait_idle()
    assert list(settings.absolute_storage_path.glob("*.original.docx"))
    assert (await client.delete(f"/api/contracts/{cid}?permanent=true")).status_code == 204
    assert not list(settings.absolute_storage_path.glob("*.original.docx")) and not list(settings.absolute_storage_path.glob("*.pdf"))


@pytest.mark.parametrize("case,name,fragment", [
    ("pdf-as-docx", "a.docx", "not a valid Word"),
    ("docx-as-pdf", "a.pdf", "missing %PDF- header"),
    ("old-doc", "a.docx", "old Word file"),
    ("broken-zip", "a.docx", "could not be opened"),
    ("plain-text", "a.docx", "not a valid Word"),
])
async def test_files_that_are_not_what_their_name_says_are_refused(client, case, name, fragment):
    payload = {
        "pdf-as-docx": b"%PDF-1.7 not really a docx", "docx-as-pdf": make_docx(),
        "old-doc": bytes.fromhex("d0cf11e0a1b11ae1") + b" old word", "broken-zip": bytes.fromhex("504b0304") + b" broken", "plain-text": b"plain text",
    }[case]
    r = await client.post("/api/contracts/upload", files={"file": (name, payload, "application/octet-stream")})
    assert r.status_code == 400 and fragment in r.json()["detail"]
    assert not list(settings.absolute_storage_path.glob("*"))       # nothing was stored


async def test_a_zip_that_is_not_a_word_document_is_refused(client):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "hi")
    r = await upload_docx(client, buf.getvalue())
    assert r.status_code == 400 and "not a Word document" in r.json()["detail"]


async def test_macro_content_is_refused(client):
    r = await upload_docx(client, make_docx(extra_zip={"word/vbaProject.bin": b"\x00macro"}))
    assert r.status_code == 400 and "macros" in r.json()["detail"]


async def test_a_zip_bomb_is_refused_before_anything_is_expanded(client):
    bomb = make_docx(extra_zip={"word/media/huge.bin": b"\x00" * (130 * 1024 * 1024)})
    assert len(bomb) < 5 * 1024 * 1024                            # tiny on the wire, huge when expanded
    r = await upload_docx(client, bomb)
    assert r.status_code == 400 and "unreasonable size" in r.json()["detail"]


async def test_a_word_document_without_text_is_refused(client):
    import docx
    buf = io.BytesIO()
    docx.Document().save(buf)
    r = await upload_docx(client, buf.getvalue())
    assert r.status_code == 400 and "no readable text" in r.json()["detail"]


async def test_word_uploads_are_owner_only(client, new_client, monkeypatch, llm):
    monkeypatch.setattr(settings, "AUTH_MODE", "login")
    a, b = client, new_client()
    for c, e in ((a, "a@x.com"), (b, "b@x.com")):
        assert (await c.post("/api/auth/register", json={"email": e, "name": "T", "password": "correct horse battery"})).status_code == 201
    cid = (await upload_docx(a, make_docx())).json()["id"]
    await job_runner.wait_idle()
    for path in (f"/api/contracts/{cid}/file", f"/api/contracts/{cid}/file?original=true"):
        assert (await a.get(path)).status_code == 200 and (await b.get(path)).status_code == 404


# ─── scanned pages ────────────────────────────────────────────────────────────

@pytest.fixture
def scanned(long_pdf):
    return to_scan(long_pdf)


@pytest.fixture
def fake_ocr(monkeypatch, long_pdf):
    """Deterministic OCR: it 'reads' each page's real text, so the pipeline can be tested without a slow model."""
    truth = page_texts(long_pdf)
    calls = {}

    async def run(pdf_path, pages, on_progress=None):
        pages = sorted(pages)
        calls["pages"] = pages
        keep = pages[: settings.OCR_MAX_PAGES]
        return ({n: PageOcr(n, truth[n], 0.95 if n != 3 else 0.7, len(truth[n].splitlines())) for n in keep}, pages[settings.OCR_MAX_PAGES:])

    monkeypatch.setattr(ocr_service, "ocr_pages", run)
    return calls


async def test_a_scanned_contract_is_read_but_nothing_from_it_is_marked_verified(client, llm, scanned, fake_ocr):
    d = await upload_and_process(client, scanned, "scan.pdf")
    assert d["status"] == "READY" and d["extraction_status"] == "complete"
    assert fake_ocr["pages"] == [1, 2, 3, 4, 5]
    src = d["extraction_meta"]["source"]
    assert src["kind"] == "ocr" and src["pages"] == [1, 2, 3, 4, 5] and src["low_confidence_pages"] == [3] and src["demoted"] > 0
    assert d["fields"] and not [f for f in d["fields"] if f["status"] == "verified"]
    exp = next(f for f in d["fields"] if f["field_key"] == "expiration_date")
    assert exp["display_value"] == "2028-12-31" and exp["status"] == "needs_review" and "OCR" in exp["notes"]["ocr"]
    assert (exp["confidence"] or 0) <= 0.7
    assert not [p for p in d["parties"] if p["status"] == "verified"]
    assert not [o for o in d["obligations"] if o["verification_status"] == "verified"]
    assert d["deadlines"]                                          # dates are still computed, from unverified values


async def test_a_scan_carries_a_review_flag_naming_the_pages_and_the_low_confidence_ones(client, llm, scanned, fake_ocr):
    d = await upload_and_process(client, scanned, "scan.pdf")
    flags = (await client.get(f"/api/contracts/{d['id']}/flags")).json()
    f = next(x for x in flags if x["code"] == "SCANNED_PAGES")
    assert f["severity"] == "HIGH" and "OCR" in f["reason"] and "Low reading confidence on page(s) 3" in f["reason"]
    assert "1, 2, 3, 4, 5" in f["reason"]


async def test_a_person_can_confirm_an_ocr_value_after_checking_the_image(client, llm, scanned, fake_ocr):
    d = await upload_and_process(client, scanned, "scan.pdf")
    exp = next(f for f in d["fields"] if f["field_key"] == "expiration_date")
    r = await client.post(f"/api/contracts/{d['id']}/fields/{exp['id']}/review", json={"action": "confirm"})
    assert r.status_code == 200 and r.json()["status"] == "verified" and r.json()["notes"]["user_review"]["action"] == "confirmed"


async def test_pages_beyond_the_ocr_limit_are_reported_as_an_incomplete_analysis(client, llm, scanned, fake_ocr, monkeypatch):
    monkeypatch.setattr(settings, "OCR_MAX_PAGES", 2)
    d = await upload_and_process(client, scanned, "scan.pdf")
    src = d["extraction_meta"]["source"]
    assert src["pages"] == [1, 2] and src["skipped_pages"] == [3, 4, 5]
    flags = (await client.get(f"/api/contracts/{d['id']}/flags")).json()
    f = next(x for x in flags if x["code"] == "SCANNED_PAGES")
    assert "3 further page(s) were not read" in f["reason"] and "incomplete" in f["reason"]


async def test_without_ocr_a_scan_is_reported_as_unsupported_with_a_reason(client, llm, scanned, monkeypatch):
    monkeypatch.setattr(ocr_service, "available", lambda: False)
    d = await upload_and_process(client, scanned, "scan.pdf")
    assert d["status"] == "UNSUPPORTED" and "no text layer" in d["extraction_error"].lower()


async def test_a_scan_with_nothing_readable_says_so(client, llm, scanned, monkeypatch):
    async def blank(pdf_path, pages, on_progress=None):
        return ({n: PageOcr(n, "", None, 0) for n in pages}, [])

    monkeypatch.setattr(ocr_service, "ocr_pages", blank)
    d = await upload_and_process(client, scanned, "scan.pdf")
    assert d["status"] == "UNSUPPORTED" and "even with OCR" in d["extraction_error"]


async def test_only_the_scanned_pages_of_a_mixed_document_are_ocred_and_demoted(client, llm, long_pdf, monkeypatch):
    src = pymupdf.open(stream=long_pdf, filetype="pdf")
    mixed = pymupdf.open()
    for i, page in enumerate(src):
        if i == 4:                                              # the last page becomes an image
            png = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False).tobytes("png")
            np = mixed.new_page(width=page.rect.width, height=page.rect.height)
            np.insert_image(np.rect, stream=png)
        else:
            mixed.insert_pdf(src, from_page=i, to_page=i)
    truth = page_texts(long_pdf)
    seen = {}

    async def run(pdf_path, pages, on_progress=None):
        seen["pages"] = sorted(pages)
        return ({n: PageOcr(n, truth[n], 0.95, 10) for n in pages}, [])

    monkeypatch.setattr(ocr_service, "ocr_pages", run)
    d = await upload_and_process(client, mixed.tobytes(), "mixed.pdf")
    assert seen["pages"] == [5] and d["extraction_meta"]["source"]["pages"] == [5]
    on_text_pages = [f for f in d["fields"] if f["page"] and f["page"] < 5]
    assert on_text_pages and all(f["status"] == "verified" for f in on_text_pages if f["status"] != "not_found")


async def test_a_reprocessed_scan_keeps_a_persons_confirmation(client, llm, scanned, fake_ocr):
    d = await upload_and_process(client, scanned, "scan.pdf")
    exp = next(f for f in d["fields"] if f["field_key"] == "expiration_date")
    await client.post(f"/api/contracts/{d['id']}/fields/{exp['id']}/review", json={"action": "confirm"})
    await client.post(f"/api/contracts/{d['id']}/reprocess")
    await job_runner.wait_idle()
    d2 = (await client.get(f"/api/contracts/{d['id']}")).json()
    assert next(f for f in d2["fields"] if f["field_key"] == "expiration_date")["notes"]["user_review"]["action"] == "confirmed"


# ─── the OCR engine itself ────────────────────────────────────────────────────

def test_lines_are_rebuilt_in_reading_order():
    box = lambda x, y, t, c=0.9: ([[x, y], [x + 80, y], [x + 80, y + 20], [x, y + 20]], t, c)
    res = [box(200, 102, "Fee:"), box(40, 100, "Monthly"), box(40, 160, "Second line"), box(120, 101, "fee")]
    text, conf = ocr_service._reading_order(res)
    assert text.splitlines() == ["Monthly fee Fee:", "Second line"] and conf == pytest.approx(0.9)
    assert ocr_service._reading_order([]) == ("", None)
    assert ocr_service._reading_order([([[0, 0], [1, 0], [1, 1], [0, 1]], "  ", 0.5)]) == ("", None)


def test_only_verified_items_on_ocr_pages_are_demoted():
    rows = [
        {"page": 2, "status": "verified", "confidence": 0.95, "notes": None},
        {"page": 1, "status": "verified", "confidence": 0.95, "notes": None},
        {"page": 2, "status": "needs_review", "confidence": 0.5, "notes": None},
        {"source_page": 2, "verification_status": "verified", "confidence": 0.9},
        {"page": 2, "status": "not_found", "confidence": None, "notes": None},
    ]
    assert ocr_service.demote(rows, [2]) == 2
    assert [r.get("status") or r.get("verification_status") for r in rows] == ["needs_review", "verified", "needs_review", "needs_review", "not_found"]
    assert rows[0]["confidence"] == 0.7 and "OCR" in rows[0]["notes"]["ocr"] and rows[1]["confidence"] == 0.95


def test_a_real_page_image_is_read_by_the_real_ocr_engine():
    if not ocr_service.available():
        pytest.skip("OCR engine not installed")
    doc = pymupdf.open()
    page = doc.new_page(width=420, height=200)
    page.insert_text((30, 60), "Effective Date: January 15, 2026", fontsize=16)
    page.insert_text((30, 110), "Monthly fee: $12,000", fontsize=16)
    png = page.get_pixmap(matrix=pymupdf.Matrix(3, 3), alpha=False).tobytes("png")
    text, conf = ocr_service.ocr_image(png)
    assert "January 15, 2026" in text and "12,000" in text and conf and conf > 0.8
